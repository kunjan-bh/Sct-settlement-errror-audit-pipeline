"""
Excel export of whatever the Disputes page is currently showing.

This is the sheet that gets sent to an aggregator to verify a settlement, so
it leads with the fields they need to find the transaction on their side --
MID, CRRN, amount, date -- and carries our own decision and notes after them.

Styled to match the session and settlement reports, so the three do not look
like they came from three different systems.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEAD_FILL = PatternFill("solid", fgColor="1F2937")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, size=14, color="111827")
_LABEL_FONT = Font(bold=True, size=10, color="374151")
_RISK_FILL = PatternFill("solid", fgColor="FEF3C7")
_NEG_FILL = PatternFill("solid", fgColor="FEE2E2")
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# (header, key, width, number format). Ordered for the person receiving it:
# what identifies the transaction first, then why it failed, then what we did.
_COLUMNS: list[tuple[str, str, int, str | None]] = [
    ("MID", "mid", 20, None),
    ("Merchant", "merchant_name", 26, None),
    ("CRRN", "crrn", 16, None),
    ("STAN", "stan", 10, None),
    ("Amount", "amount", 14, "#,##0.00"),
    ("Date / time", "date_time", 19, None),
    ("Aggregator / wallet", "mapped_partner", 20, None),
    ("Institution (switch)", "bank_or_wallet", 22, None),
    ("Reason", "reason", 40, None),
    ("Stopped at", "current_status", 22, None),
    ("Hold balance", "hold_balance", 14, "#,##0.00"),
    ("Total balance", "total_balance", 14, "#,##0.00"),
    ("Verify before retry", "_risk", 18, None),
    ("Our status", "_status", 14, None),
    ("Note", "op_comment", 40, None),
    ("Creditor account", "creditor_account", 22, None),
    ("Ref ID", "ref_id", 20, None),
]

_STATUS_LABEL = {
    "pending": "Open",
    "in_progress": "In Progress",
    "solved": "Solved",
    "exclude": "Excluded",
}


def _cell_value(row: dict, key: str):
    if key == "_risk":
        return "YES" if row.get("double_pay_risk") else ""
    if key == "_status":
        return _STATUS_LABEL.get(row.get("op_status") or "pending", row.get("op_status"))
    if key == "date_time":
        return str(row.get("date_time") or "")[:19].replace("T", " ")
    return row.get(key)


def generate_dispute_xlsx(rows: list[dict], meta: dict) -> bytes:
    """
    One sheet of disputes, plus a short header saying what this is a list of.

    `meta` carries the range and the filter that was applied, because a sheet
    of 12 rows with no context is impossible to interpret a week later -- the
    reader cannot tell whether it is everything or a slice.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Disputes"

    ws["A1"] = "Settlement disputes"
    ws["A1"].font = _TITLE_FONT

    header_bits = [
        ("Range", f"{meta.get('date_from', '')} to {meta.get('date_to', '')}"),
        ("Showing", meta.get("filter_label") or "all"),
        ("Rows", len(rows)),
        ("Total amount", round(sum(float(r.get("amount") or 0) for r in rows), 2)),
        ("Generated", meta.get("generated_at", "")),
    ]
    r = 2
    for label, value in header_bits:
        ws.cell(row=r, column=1, value=label).font = _LABEL_FONT
        cell = ws.cell(row=r, column=2, value=value)
        if label == "Total amount":
            cell.number_format = "#,##0.00"
        r += 1

    head_row = r + 1
    for col, (label, _key, width, _fmt) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=head_row, column=col, value=label)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = width

    for i, row in enumerate(rows):
        excel_row = head_row + 1 + i
        # Tint the rows that need a decision from someone else: a connection
        # drop that could pay twice, or a merchant whose ledger is broken.
        tint = (
            _NEG_FILL if row.get("negative_hold")
            else _RISK_FILL if row.get("double_pay_risk")
            else None
        )
        for col, (_label, key, _w, fmt) in enumerate(_COLUMNS, start=1):
            cell = ws.cell(row=excel_row, column=col, value=_cell_value(row, key))
            cell.border = _BORDER
            if fmt:
                cell.number_format = fmt
            if key in ("reason", "op_comment"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if tint:
                cell.fill = tint

    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{head_row}:{get_column_letter(len(_COLUMNS))}{head_row + len(rows)}"
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
