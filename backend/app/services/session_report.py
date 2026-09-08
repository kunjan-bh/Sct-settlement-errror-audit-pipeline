"""
The report for one office session.

A session has no ingested spreadsheet behind it, so the five transaction
sheets the upload report builds have nothing to draw from. What a session has
is a list of decisions: which failed settlements were looked at, what was
concluded, and what was written down. That is what this workbook is.

It deliberately mirrors the upload report's shape -- an Overview sheet, then
detail -- so the two kinds of batch produce something recognisably the same
when they land in the same inbox.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.batch import Batch
from app.services.session_service import session_activity

_HEAD_FILL = PatternFill("solid", fgColor="1F2937")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, size=14, color="111827")
_LABEL_FONT = Font(bold=True, size=10, color="374151")
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COLUMNS = [
    ("MID", "mid", 20),
    ("CRRN", "crrn", 18),
    ("Amount", "amount", 14),
    ("Aggregator / wallet", "partner_name", 24),
    ("Decision", "status", 14),
    ("Comment", "comment", 52),
    ("Decided at", "updated_at", 20),
]


def _write_header(ws, row: int) -> int:
    for col, (label, _key, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    return row + 1


def _write_rows(ws, row: int, records: list[dict]) -> int:
    for rec in records:
        for col, (_label, key, _w) in enumerate(_COLUMNS, start=1):
            value = rec.get(key)
            if key == "status":
                value = "In Progress" if value == "in_progress" else str(value or "").title()
            elif key == "updated_at" and value:
                value = str(value)[:19].replace("T", " ")
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = _BORDER
            if key == "amount":
                cell.number_format = "#,##0.00"
            elif key == "comment":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    return row


def generate_session_report_bytes(batch_id: int) -> bytes:
    """The session's workbook: an overview, then the settlements handled."""
    batch = Batch.query.get_or_404(batch_id)
    activity = session_activity(batch_id)
    totals = activity["totals"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Session Overview"

    ws["A1"] = batch.name
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:C1")

    rows = [
        ("Session", batch.name),
        ("Opened", str(batch.created_at)[:19].replace("T", " ") if batch.created_at else "—"),
        ("Closed", str(batch.finished_at)[:19].replace("T", " ") if batch.finished_at else "still open"),
        ("", ""),
        ("Settlements handled", totals["handled"]),
        ("Solved", totals["solved"]),
        ("In progress", totals["in_progress"]),
        ("Amount handled", totals["amount_handled"]),
        ("Amount solved", totals["amount_solved"]),
    ]
    r = 3
    for label, value in rows:
        if not label:
            r += 1
            continue
        ws.cell(row=r, column=1, value=label).font = _LABEL_FONT
        cell = ws.cell(row=r, column=2, value=value)
        if "Amount" in label:
            cell.number_format = "#,##0.00"
        r += 1

    if batch.notes:
        r += 1
        ws.cell(row=r, column=1, value="Notes").font = _LABEL_FONT
        note = ws.cell(row=r, column=2, value=batch.notes)
        note.alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60

    # Excluded decisions are not written to any sheet. An exclusion is not work
    # done -- it is work judged not ours -- and putting it in a report of what
    # was handled would overstate the day.
    for title, records in (
        ("Solved", activity["solved"]),
        ("In Progress", activity["in_progress"]),
    ):
        sheet = wb.create_sheet(title)
        sheet["A1"] = f"{title} — {len(records)} settlement{'' if len(records) == 1 else 's'}"
        sheet["A1"].font = _TITLE_FONT
        start = _write_header(sheet, 3)
        _write_rows(sheet, start, records)
        sheet.freeze_panes = sheet.cell(row=start, column=1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
